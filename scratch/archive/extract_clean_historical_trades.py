"""
extract_clean_historical_trades.py
==================================
Milestone M1 — Comprehensive Historical Trade Extraction & Clean Data Pipeline.

Extracts, normalizes, and deterministically classifies 100% of historical trade records
from both SQLite master storage (output/monitor/trades.sqlite3) and the pre-migration Excel
archive (output/exports/trade_archive.xlsx). Also consolidates real-time candidate scans
from output/monitor/trade_journal.csv.

Primary Deliverables:
- output/extracted_trades_clean.json : Complete structured dataset with metadata, trades, and candidate scans.
- output/extracted_trades_clean.csv  : Tabular CSV of all 1,097 historical trades with normalized columns.

Deterministic Invariants Enforced:
- SQLite Historical Snapshot (IDs 571..1133): Exactly 536 rows.
- Excel Historical Archive (July/August 2026): Exactly 561 rows.
- Total Unique Records: Exactly 1,097.
- Synthetic Test Fixtures: Exactly 168 (166 unit_test_risk_engine + 2 test).
- Genuine Market Trades: Exactly 929 (368 in SQLite + 561 in XLSX).
- Filled Trades: 607 (179 in SQLite + 428 in XLSX).
- Unfilled/Cancelled Limit Orders: 490 (189 SQLite genuine + 133 XLSX genuine + 168 synthetic).
- CLOSED_EXTERNALLY Unfilled Limit Orders (order_status == 'OPEN'): Exactly 110.
"""

import os
import sys
import json
import sqlite3
import re
import argparse
import logging
from datetime import datetime
from collections import Counter
import pandas as pd
import openpyxl

# Setup Project Paths
PROJ_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJ_ROOT not in sys.path:
    sys.path.insert(0, PROJ_ROOT)
COMMON_DIR = os.path.join(PROJ_ROOT, "common")
if COMMON_DIR not in sys.path:
    sys.path.insert(0, COMMON_DIR)

try:
    import paths
    OUTPUT_DIR = os.path.dirname(paths.MONITOR_DIR)
    MONITOR_DIR = paths.MONITOR_DIR
    EXPORTS_DIR = paths.EXPORTS_DIR
    DB_PATH = os.path.join(MONITOR_DIR, "trades.sqlite3")
    XLSX_PATH = paths.TRADE_ARCHIVE_XLSX
    JOURNAL_PATH = paths.TRADE_JOURNAL_CSV
except ImportError:
    OUTPUT_DIR = os.path.join(PROJ_ROOT, "output")
    MONITOR_DIR = os.path.join(OUTPUT_DIR, "monitor")
    EXPORTS_DIR = os.path.join(OUTPUT_DIR, "exports")
    DB_PATH = os.path.join(MONITOR_DIR, "trades.sqlite3")
    XLSX_PATH = os.path.join(EXPORTS_DIR, "trade_archive.xlsx")
    JOURNAL_PATH = os.path.join(MONITOR_DIR, "trade_journal.csv")

OUT_JSON = os.path.join(OUTPUT_DIR, "extracted_trades_clean.json")
OUT_CSV = os.path.join(OUTPUT_DIR, "extracted_trades_clean.csv")

# Baseline cutoff for SQLite historical master database (frozen at start of M1 audit)
MAX_HISTORICAL_DB_ID = 1133

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("extract_clean_trades")


def parse_contract_details(contract_str):
    """
    Parse derivative contract string to extract underlying symbol, strike price, and option side.
    Handles Indian F&O naming conventions:
    - Monthly: <SYMBOL><YY><MONTH_3_LETTER><STRIKE><CE/PE> (e.g. BEL26AUG390CE, NIFTY26SEP25450PE)
    - Weekly:  <SYMBOL><YY><M><DD><STRIKE><CE/PE> (e.g. SENSEX2681377800PE -> SENSEX, 77800, PE)
    - Fallback generic strike extraction
    """
    if not contract_str:
        return None, None, None
    c = str(contract_str).strip().upper()
    if ":" in c:
        c = c.split(":")[-1]

    # Monthly contracts (e.g. BEL26AUG390CE, INFY26SEP1800PE)
    m_month = re.match(r"^([A-Z&-]+?)\d{2}[A-Z]{3}(\d+(?:\.\d+)?)(CE|PE)$", c)
    if m_month:
        return m_month.group(1), float(m_month.group(2)), m_month.group(3)

    # Weekly index contracts (e.g. SENSEX2681377800PE -> Year 26, Month 8, Day 13, Strike 77800 PE)
    m_week = re.match(r"^([A-Z&-]+?)\d{2}[1-9OND]\d{2}(\d+(?:\.\d+)?)(CE|PE)$", c)
    if m_week:
        return m_week.group(1), float(m_week.group(2)), m_week.group(3)

    # Generic fallback
    m_gen = re.search(r"(\d+(?:\.\d+)?)(CE|PE)$", c)
    if m_gen:
        return None, float(m_gen.group(1)), m_gen.group(2)

    return None, None, None


def is_synthetic(engine, symbol, contract, data):
    """
    Deterministically classify whether a record is a synthetic/test fixture vs genuine trade.
    Enforces the four classification rules from worker_m1_extractor_spec.md:
    1. Engine rule (unit_test_risk_engine, test, mock, dummy)
    2. Symbol / contract test markers (GHOST, TEST, FAIL, MOCK, DUMMY, etc.)
    3. Mock order ID rule (non-15-digit non-Kite ID)
    4. Explicit test flags in data payload
    """
    eng = (engine or "").lower()
    if engine in ("unit_test_risk_engine", "test") or any(k in eng for k in ("test", "mock", "dummy", "reg_test")):
        return True, f"Engine test fixture: {engine}"

    s_upper = (symbol or "").upper()
    c_upper = (contract or "").upper()
    test_markers = [
        "GHOST", "TEST", "FAIL", "MOCK", "DUMMY",
        "FOMO_SYM", "BOUND_LOWER", "BOUND_UPPER", "ZQ_FAIL", "IDX_FAIL"
    ]
    for marker in test_markers:
        if marker in s_upper or marker in c_upper:
            return True, f"Symbol/Contract test marker: {marker}"

    order_id = str(data.get("order_id") or "")
    if order_id and order_id not in ("None", "", "null"):
        # Zerodha Kite order IDs are exactly 15 digits
        if not re.match(r"^\d{15}$", order_id):
            return True, f"Mock Order ID: {order_id}"

    if data.get("is_test") or data.get("mock") or data.get("is_mock"):
        return True, "Explicit test flag in data payload"

    return False, None


def extract_sqlite_trades(db_path, max_id=MAX_HISTORICAL_DB_ID):
    """
    Extract all historical rows from SQLite master database up to max_id.
    Normalizes metadata and parses JSON data payload.
    """
    if not os.path.exists(db_path):
        raise FileNotFoundError(f"SQLite database not found at {db_path}")

    conn = sqlite3.connect(db_path)
    cur = conn.cursor()

    # Check total available rows in DB
    cur.execute("SELECT count(*) FROM trades")
    total_in_db = cur.fetchone()[0]

    query = (
        "SELECT id, engine, symbol, contract, status, created_at, updated_at, data_json "
        "FROM trades "
    )
    params = []
    if max_id is not None:
        query += "WHERE id <= ? ORDER BY id"
        params.append(max_id)
    else:
        query += "ORDER BY id"

    cur.execute(query, params)
    rows = cur.fetchall()
    conn.close()

    logger.info(
        f"SQLite extraction: retrieved {len(rows)} rows (max_id={max_id}) "
        f"from total {total_in_db} in {db_path}."
    )

    trades = []
    for row in rows:
        tid, eng, sym, cnt, st, cat, uat, dj = row
        d = json.loads(dj) if dj else {}

        synth, synth_reason = is_synthetic(eng, sym, cnt, d)

        # Parse contract metadata for missing fields
        parsed_sym, parsed_strike, parsed_side = parse_contract_details(cnt)
        strike = float(d["strike"]) if d.get("strike") is not None else parsed_strike
        side = d.get("side") or parsed_side
        direction = d.get("direction") or (
            "BULL" if side in ("CE", "BUY") else ("BEAR" if side in ("PE", "SELL") else "BULL")
        )

        entry_price = float(d["entry_price"]) if d.get("entry_price") is not None else None
        exit_price = float(d["exit_price"]) if d.get("exit_price") is not None else None
        entry_spot = float(d["entry_spot"]) if d.get("entry_spot") is not None else (
            float(d["spot_entry"]) if d.get("spot_entry") is not None else None
        )
        spot_anchor = float(d["anchor_floor"]) if d.get("anchor_floor") is not None else entry_spot
        pnl_pct = float(d["pnl_percent"]) if d.get("pnl_percent") is not None else None
        pnl_amt = float(d["pnl_amount"]) if d.get("pnl_amount") is not None else None

        # Determine filled invariant
        has_fill = False
        if not synth:
            has_fill = bool(
                d.get("order_status") == "FILLED"
                or (entry_price is not None and entry_price > 0)
                or st in ("SL_HIT", "TARGET_HIT", "USER_EXIT")
                or (st == "COMPLETED" and pnl_pct is not None and pnl_pct != 0)
            )

        # Extract Point D RVOL from point_d_rvol or trade_dna.spot_rvol
        dna = d.get("trade_dna") or {}
        point_d_rvol = None
        if d.get("point_d_rvol") is not None:
            point_d_rvol = float(d["point_d_rvol"])
        elif isinstance(dna, dict) and dna.get("spot_rvol") is not None:
            point_d_rvol = float(dna["spot_rvol"])

        sl_val = float(d["current_sl"]) if d.get("current_sl") is not None else (
            float(d["sl"]) if d.get("sl") is not None else None
        )
        spot_sl_val = float(d["spot_sl"]) if d.get("spot_sl") is not None else (
            float(d["geometric_sl"]) if d.get("geometric_sl") is not None else None
        )

        entry = {
            "trade_id": f"DB_{tid}",
            "source_id": tid,
            "data_source": "trades_sqlite",
            "engine": eng,
            "symbol": sym,
            "contract": cnt or d.get("contract"),
            "timestamp": cat,
            "pattern": d.get("pattern") or "NONE",
            "strike": strike,
            "side": side,
            "direction": direction,
            "timeframe": d.get("timeframe"),
            "tier": d.get("tier_label") or d.get("tier") or d.get("tier_badge"),
            "execution_type": d.get("execution_type") or ("TEST" if synth else "ALGO_TRIGGER"),
            "entry_price": entry_price,
            "exit_price": exit_price,
            "entry_spot": entry_spot,
            "spot_anchor": spot_anchor,
            "sl": sl_val,
            "spot_sl": spot_sl_val,
            "t1": float(d["t1"]) if d.get("t1") is not None else None,
            "t2": float(d["t2"]) if d.get("t2") is not None else None,
            "t3": float(d["t3"]) if d.get("t3") is not None else None,
            "rr": float(d["rr"]) if d.get("rr") is not None else None,
            "realized_pnl_pct": pnl_pct,
            "realized_pnl_amt": pnl_amt,
            "original_status": st,
            "order_status": d.get("order_status"),
            "order_id": str(d.get("order_id")) if d.get("order_id") is not None else None,
            "exit_reason": d.get("exit_reason"),
            "details": d.get("details"),
            "is_synthetic": synth,
            "synthetic_reason": synth_reason,
            "is_filled": has_fill,
            "point_d_rvol": point_d_rvol,
            "volume": float(d["volume"]) if d.get("volume") is not None else None
        }
        trades.append(entry)

    return trades


def extract_xlsx_trades(xlsx_path):
    """
    Extract historical trades from output/exports/trade_archive.xlsx across all sheets.
    Normalizes columns to align with SQLite schema.
    """
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Trade archive Excel workbook not found at {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path)
    trades = []

    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        rows = list(ws.iter_rows(values_only=True))[1:]  # skip header
        sheet_count = 0
        for row in rows:
            tid = row[0]
            if tid is None:
                continue

            cnt = str(row[17]) if len(row) > 17 and row[17] else None
            side_val = str(row[16]) if len(row) > 16 and row[16] else None
            parsed_sym, parsed_strike, parsed_side = parse_contract_details(cnt)
            side = side_val or parsed_side
            strike = parsed_strike
            direction = "BULL" if side == "CE" else ("BEAR" if side == "PE" else "BULL")

            status_str = str(row[4]) if row[4] is not None else "UNKNOWN"
            has_fill = status_str in ("COMPLETED", "SL_HIT", "TARGET_HIT", "USER_EXIT")
            entry_spot = float(row[5]) if row[5] is not None else None
            sl_val = float(row[6]) if row[6] is not None else None
            t1_val = float(row[7]) if row[7] is not None else None
            t2_val = float(row[8]) if row[8] is not None else None
            t3_val = float(row[9]) if row[9] is not None else None

            rr_val = None
            if entry_spot and sl_val and t1_val and (entry_spot - sl_val) != 0:
                rr_val = round((t1_val - entry_spot) / (entry_spot - sl_val), 2)

            pnl_pct = float(row[15]) if row[15] is not None else None

            entry = {
                "trade_id": f"XLSX_{tid}",
                "source_id": int(tid),
                "data_source": "trade_archive_xlsx",
                "engine": str(row[1]) if row[1] is not None else "nifty50",
                "symbol": str(row[2]) if row[2] is not None else None,
                "contract": cnt,
                "timestamp": str(row[13]) if row[13] is not None else None,
                "pattern": str(row[3]) if row[3] is not None else "NONE",
                "strike": strike,
                "side": side,
                "direction": direction,
                "timeframe": None,
                "tier": None,
                "execution_type": "HISTORICAL_ARCHIVE",
                "entry_price": entry_spot,
                "exit_price": None,
                "entry_spot": entry_spot,
                "spot_anchor": entry_spot,
                "sl": sl_val,
                "spot_sl": sl_val,
                "t1": t1_val,
                "t2": t2_val,
                "t3": t3_val,
                "rr": rr_val,
                "realized_pnl_pct": pnl_pct,
                "realized_pnl_amt": None,
                "original_status": status_str,
                "order_status": "FILLED" if has_fill else "FAILED",
                "order_id": None,
                "exit_reason": status_str if status_str in ("SL_HIT", "TARGET_HIT", "USER_EXIT") else None,
                "details": f"Archived {sheetname} trade",
                "is_synthetic": False,
                "synthetic_reason": None,
                "is_filled": has_fill,
                "point_d_rvol": None,
                "volume": None
            }
            trades.append(entry)
            sheet_count += 1
        logger.info(f"Excel extraction: {sheetname} yielded {sheet_count} rows.")

    logger.info(f"Excel extraction total: {len(trades)} trades across {len(wb.sheetnames)} sheets.")
    return trades


def extract_candidate_scans(journal_csv_path):
    """
    Extract candidate scan setups from output/monitor/trade_journal.csv.
    Filters on SCAN_MATCH events and deduplicates by (timestamp, contract, pattern).
    """
    if not os.path.exists(journal_csv_path):
        logger.warning(f"Journal CSV not found at {journal_csv_path}, skipping candidate scans.")
        return []

    import csv
    candidates = []
    seen = set()

    with open(journal_csv_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f, delimiter="\t")
        idx = 1
        for row in reader:
            if len(row) >= 11 and row[4] == "SCAN_MATCH":
                ts = row[0]
                contract = row[1]
                pattern = row[2]
                key = (ts, contract, pattern)
                if key in seen:
                    continue
                seen.add(key)

                # Extract tier and side from details or contract
                details_str = row[10]
                tier_match = re.search(r"Tier=([A-Z0-9_]+)", details_str)
                tier = tier_match.group(1) if tier_match else None

                side_match = re.search(r"Side=([A-Z]+)", details_str)
                side = side_match.group(1) if side_match else (
                    "CE" if contract.endswith("CE") else ("PE" if contract.endswith("PE") else None)
                )

                # Symbol extraction
                sym_match = re.match(r"^([A-Z&-]+?)\d", contract)
                symbol = sym_match.group(1) if sym_match else contract

                def to_float(val):
                    try:
                        return float(val) if val not in ("", "-", "None", None) else None
                    except (ValueError, TypeError):
                        return None

                cand = {
                    "candidate_id": f"SCAN_{idx}",
                    "source": "trade_journal_csv",
                    "timestamp": ts,
                    "symbol": symbol,
                    "contract": contract,
                    "pattern": pattern,
                    "timeframe": row[3],
                    "side": side,
                    "entry": to_float(row[6]),
                    "sl": to_float(row[7]),
                    "target": to_float(row[8]),
                    "rr": to_float(row[9]),
                    "tier": tier,
                    "status": row[5],
                    "details": details_str
                }
                candidates.append(cand)
                idx += 1

    logger.info(f"Candidate scans extraction: parsed {len(candidates)} unique SCAN_MATCH setups.")
    return candidates


def generate_clean_dataset(max_db_id=MAX_HISTORICAL_DB_ID):
    """
    Main pipeline execution:
    1. Extract SQLite rows (id <= max_db_id)
    2. Extract XLSX rows
    3. Extract candidate scans
    4. Validate invariants
    5. Save JSON and CSV
    """
    logger.info("=" * 60)
    logger.info("STARTING HISTORICAL TRADE EXTRACTION & CLEAN PIPELINE")
    logger.info("=" * 60)

    db_trades = extract_sqlite_trades(DB_PATH, max_id=max_db_id)
    xlsx_trades = extract_xlsx_trades(XLSX_PATH)
    candidate_scans = extract_candidate_scans(JOURNAL_PATH)

    all_trades = db_trades + xlsx_trades

    # Metrics computation
    total_records = len(all_trades)
    synthetic_records = [t for t in all_trades if t["is_synthetic"]]
    genuine_records = [t for t in all_trades if not t["is_synthetic"]]

    db_genuine = [t for t in db_trades if not t["is_synthetic"]]
    db_synthetic = [t for t in db_trades if t["is_synthetic"]]
    xlsx_genuine = [t for t in xlsx_trades if not t["is_synthetic"]]
    xlsx_synthetic = [t for t in xlsx_trades if t["is_synthetic"]]

    filled_trades = [t for t in all_trades if t["is_filled"]]
    unfilled_trades = [t for t in all_trades if not t["is_filled"]]
    genuine_filled = [t for t in genuine_records if t["is_filled"]]
    genuine_unfilled = [t for t in genuine_records if not t["is_filled"]]

    # CLOSED_EXTERNALLY open order forensic verification
    closed_ext_open = [
        t for t in db_genuine
        if t["original_status"] == "CLOSED_EXTERNALLY" and t["order_status"] == "OPEN"
    ]

    pattern_counts = Counter(t["pattern"] for t in all_trades)
    genuine_pattern_counts = Counter(t["pattern"] for t in genuine_records)
    engine_counts = Counter(t["engine"] for t in all_trades)
    status_counts = Counter(t["original_status"] for t in all_trades)

    metadata = {
        "extracted_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_records": total_records,
        "genuine_trades_count": len(genuine_records),
        "synthetic_fixtures_count": len(synthetic_records),
        "filled_trades_count": len(filled_trades),
        "unfilled_trades_count": len(unfilled_trades),
        "genuine_filled_count": len(genuine_filled),
        "genuine_unfilled_count": len(genuine_unfilled),
        "sources_summary": {
            "trades_sqlite": {
                "total": len(db_trades),
                "genuine": len(db_genuine),
                "synthetic": len(db_synthetic),
                "filled": sum(1 for t in db_trades if t["is_filled"]),
                "unfilled": sum(1 for t in db_trades if not t["is_filled"]),
                "closed_externally_open_count": len(closed_ext_open)
            },
            "trade_archive_xlsx": {
                "total": len(xlsx_trades),
                "genuine": len(xlsx_genuine),
                "synthetic": len(xlsx_synthetic),
                "filled": sum(1 for t in xlsx_trades if t["is_filled"]),
                "unfilled": sum(1 for t in xlsx_trades if not t["is_filled"])
            }
        },
        "candidate_scans_count": len(candidate_scans),
        "engine_distribution": dict(engine_counts),
        "status_distribution": dict(status_counts),
        "pattern_distribution_all": dict(pattern_counts),
        "pattern_distribution_genuine": dict(genuine_pattern_counts)
    }

    # Invariant Assertions
    logger.info("VERIFYING DETERMINISTIC INVARIANTS...")
    assert len(db_trades) == 536, f"Expected 536 SQLite rows, got {len(db_trades)}"
    assert len(xlsx_trades) == 561, f"Expected 561 XLSX rows, got {len(xlsx_trades)}"
    assert total_records == 1097, f"Expected 1097 total records, got {total_records}"
    assert len(synthetic_records) == 168, f"Expected 168 synthetic records, got {len(synthetic_records)}"
    assert len(genuine_records) == 929, f"Expected 929 genuine records, got {len(genuine_records)}"
    assert len(db_genuine) == 368, f"Expected 368 SQLite genuine records, got {len(db_genuine)}"
    assert len(xlsx_genuine) == 561, f"Expected 561 XLSX genuine records, got {len(xlsx_genuine)}"
    assert len(filled_trades) == 607, f"Expected 607 filled trades, got {len(filled_trades)}"
    assert len(unfilled_trades) == 490, f"Expected 490 unfilled trades, got {len(unfilled_trades)}"
    assert len(closed_ext_open) == 110, f"Expected 110 CLOSED_EXTERNALLY OPEN orders, got {len(closed_ext_open)}"
    logger.info("ALL 10 DETERMINISTIC INVARIANTS CONFIRMED 100% PASS.")

    # Save to JSON
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    payload = {
        "metadata": metadata,
        "trades": all_trades,
        "candidate_scans": candidate_scans
    }
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, default=str)
    logger.info(f"Clean JSON dataset written to {OUT_JSON} ({os.path.getsize(OUT_JSON):,} bytes)")

    # Save to CSV
    df = pd.DataFrame(all_trades)
    df.to_csv(OUT_CSV, index=False, encoding="utf-8")
    logger.info(f"Clean CSV dataset written to {OUT_CSV} ({len(df)} rows, {len(df.columns)} columns)")

    logger.info("=" * 60)
    logger.info("HISTORICAL TRADE EXTRACTION COMPLETED SUCCESSFULLY")
    logger.info(f"Total: {total_records} | Genuine: {len(genuine_records)} | Synthetic: {len(synthetic_records)}")
    logger.info(f"Filled: {len(filled_trades)} (Genuine: {len(genuine_filled)}) | Unfilled: {len(unfilled_trades)}")
    logger.info(f"Artifacts: {OUT_JSON} and {OUT_CSV}")
    logger.info("=" * 60)

    return metadata, all_trades, candidate_scans


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract clean historical trades dataset.")
    parser.add_argument(
        "--max-db-id",
        type=int,
        default=MAX_HISTORICAL_DB_ID,
        help=f"Max SQLite trade ID to include (default: {MAX_HISTORICAL_DB_ID})"
    )
    args = parser.parse_args()
    generate_clean_dataset(max_db_id=args.max_db_id)
